import sys, datetime as dt, collections
sys.path.insert(0, '/tmp/claude-0/w')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule, ColorScaleRule
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from cons import load, area_union, nstatus, pdate

SRC = 'sp.xlsx'; OUT = 'LCQ_RJ_COCKPIT_GESTAO_DOCUMENTAL_EXECUTIVO_FINAL.xlsx'
FONT = 'Segoe UI'
# ---------- paleta ----------
NAVY = '1F2A44'; BLUE = '2F5597'; BLUE2 = '8EA9D6'; BLUEL = 'E8EEF7'; BG = 'F3F5F8'; WHITE = 'FFFFFF'
TXT = '1F2937'; MUTED = '6B7280'; LINE = 'DDE2EA'
RED = 'B42318'; REDL = 'FBEAE8'; AMB = '9A6700'; AMBL = 'FDF3DC'; GRN = '1E7B4C'; GRNL = 'E6F4EC'; GRY = '8A94A3'; GRYL = 'EEF0F3'

def F(size=10, bold=False, color=TXT, italic=False, u=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=u)
def FILL(c): return PatternFill('solid', start_color=c, end_color=c)
def side(c=LINE, s='thin'): return Side(style=s, color=c)
AL = lambda h='left', v='center', w=False, ind=0: Alignment(horizontal=h, vertical=v, wrap_text=w, indent=ind)

def link(c, sheet, text=None, cell='A1'):
    if text is not None: c.value = text
    c.hyperlink = Hyperlink(ref=c.coordinate, location="'" + sheet.replace("'", "''") + f"'!{cell}", display=str(c.value))

def box(ws, rng, fill=None, font=None, align=None, border=None):
    for row in ws[rng]:
        for c in row:
            if fill: c.fill = fill
            if font: c.font = font
            if align: c.alignment = align
            if border: c.border = border

def mput(ws, rng, value, font=None, fill=None, align=None):
    ws.merge_cells(rng)
    c = ws[rng.split(':')[0]]
    c.value = value
    if font: c.font = font
    if fill: box(ws, rng, fill=fill)
    if align: c.alignment = align
    return c

# ---------- carga / consolidação ----------
_, recs, lvs, groups = load(SRC)
wb = openpyxl.load_workbook(SRC)
for n in ['DASHBOARD PREMIUM', '_AUX DASH', 'BASE CONSOLIDADA', 'QUALIDADE BASE', 'OUTROS CONTROLES']:
    del wb[n]

N = len(groups); R0 = 5; R1 = R0 + N - 1
BS = "'BASE CONSOLIDADA'"
def col(c): return f"{BS}!${c}${R0}:${c}${R1}"

# ---------- _AUX (cálculos) ----------
ax = wb.create_sheet('_AUX DASH')
AX = "'_AUX DASH'"
ax['A1'] = 'CÁLCULOS DE APOIO DO COCKPIT — não editar (exceto Data de referência)'; ax['A1'].font = F(12, True, NAVY)
ax['A3'] = 'Data de referência'; ax['B3'] = '=TODAY()'; ax['B3'].number_format = 'dd/mm/yyyy'
ax['C3'] = 'Padrão: =HOJE(). Para congelar uma data de reunião, digite a data em B3.'
ax['C3'].font = F(9, color=MUTED, italic=True)
areas = ['Todas', 'Q4', 'PE', 'PP', 'AOL', 'COMUM']
sits = ['Todas', 'Ação imediata', 'Em tratamento', 'Validar cadastro', 'Planejar revisão', 'Regular', 'Inativo']
hors = ['Todos', 'Vencida', 'Até 90 dias', '91 dias a 12 meses', 'Após 12 meses', 'Sem validade', 'Inativo']
stats = ['Todos'] + sorted({nstatus(l[0]['st']) for l in groups.values()})
resps = ['Todos'] + sorted({str(l[0]['resp']).replace('\xa0', ' ').strip().title() for l in groups.values()})
lists = [('Áreas', areas), ('Situações', sits), ('Horizontes', hors), ('Status', stats), ('Responsáveis', resps)]
LR = {}
for j, (h, vals) in enumerate(lists):
    c = 1 + j
    ax.cell(5, c, h).font = F(9, True, MUTED)
    for i, v in enumerate(vals): ax.cell(6 + i, c, v)
    LR[h] = f"{AX}!${L(c)}$6:${L(c)}${5 + len(vals)}"

def dn(name, ref): wb.defined_names[name] = DefinedName(name, attr_text=ref)
dn('DataRef', f"{AX}!$B$3")
DB = 'DASHBOARD EXECUTIVO'; DQ = "'DASHBOARD EXECUTIVO'"
dn('fArea', f"{DQ}!$B$7"); dn('fStatus', f"{DQ}!$F$7"); dn('fSit', f"{DQ}!$J$7")
dn('fResp', f"{DQ}!$N$7"); dn('fHor', f"{DQ}!$S$7"); dn('qBusca', "'PESQUISA'!$D$6")

# ---------- BASE CONSOLIDADA ----------
bs = wb.create_sheet('BASE CONSOLIDADA')
H = ['Código NDocs', 'Versão', 'Antigo Número', 'Título', 'Status NDocs', 'Validade', 'Área consolidada', 'Abrangência',
     'Responsável', 'Norma', 'Comentários', 'Origem (aba!linha)', 'Qtd. registros origem',
     'Status padronizado', 'Validade (analítica)', 'Responsável padronizado', 'Dias p/ vencer', 'Situação Gerencial',
     'Horizonte de validade', 'Mês de vencimento', 'Prioridade', 'Ação recomendada', 'Filtro dashboard', 'Score ação',
     'Chave de busca', 'Busca (linha)']
W = [22, 7, 17, 52, 15, 11, 12, 13, 32, 32, 26, 30, 9, 15, 11, 32, 9, 17, 17, 11, 9, 40, 9, 11, 30, 9]
for i, (h, w) in enumerate(zip(H, W), 1):
    c = bs.cell(4, i, h)
    analytic = i >= 14
    c.font = F(9, True, WHITE); c.fill = FILL(BLUE if analytic else NAVY)
    c.alignment = AL('left' if i not in (2, 13, 17, 21, 23) else 'center', 'center', True)
    bs.column_dimensions[L(i)].width = w
bs.row_dimensions[4].height = 32
bs['A1'] = 'BASE CONSOLIDADA DE INSTRUÇÕES — LCQ RJ'; bs['A1'].font = F(16, True, NAVY)
bs['A2'] = (f'=TEXT(COUNTA({col("A")}),"0")&" códigos únicos · consolidados de NDocs Q4, PE_PP, AOL e COMUNS · '
            f'colunas A–M: dado original da origem (sem edição) · colunas N–Z: camada analítica calculada"')
bs['A2'].font = F(9, color=MUTED)
mput(bs, 'N3:Z3', 'CAMADA ANALÍTICA (fórmulas — não editar)', F(8, True, BLUE), align=AL('left'))
mput(bs, 'A3:M3', 'DADO ORIGINAL NDocs (valores copiados da aba de origem indicada em "Origem")', F(8, True, NAVY), align=AL('left'))
for r in (1, 2): bs.row_dimensions[r].height = 24 if r == 1 else 16

for k, (code, l) in enumerate(groups.items()):
    r = R0 + k; d = l[0]
    vals = [d['cod'], d['ver'], d['ant'], d['tit'], d['st'], d['val'], area_union(l), d['abr'], d['resp'], d['norma'], d['com'],
            '; '.join(f"{x['sheet']}!{x['row']}" for x in l), len(l)]
    for i, v in enumerate(vals, 1): bs.cell(r, i, v)
    s = lambda e: f'TRIM(SUBSTITUTE({e},CHAR(160)," "))'
    t = f'RIGHT(TRIM(F{r}),10)'
    f = {
        'N': f'=IF({s(f"E{r}")}="Avaiação de uso","Avaliação de uso",IF({s(f"E{r}")}="Cancelada","Cancelado",{s(f"E{r}")}))',
        'O': f'=IF(ISNUMBER(F{r}),F{r},IFERROR(DATE(RIGHT({t},4),MID({t},4,2),LEFT({t},2)),""))',
        'P': f'=PROPER({s(f"I{r}")})',
        'Q': f'=IF(O{r}="","",O{r}-DataRef)',
        'R': (f'=IF(OR(N{r}="Cancelado",N{r}="Fora de Uso"),"Inativo",IF(OR(N{r}="Revisão",N{r}="Em Elaboração"),"Em tratamento",'
              f'IF(OR(N{r}<>"Aprovado",O{r}=""),"Validar cadastro",IF(O{r}<DataRef,"Ação imediata",'
              f'IF(O{r}<=EDATE(DataRef,12),"Planejar revisão","Regular")))))'),
        'S': (f'=IF(R{r}="Inativo","Inativo",IF(O{r}="","Sem validade",IF(O{r}<DataRef,"Vencida",IF(O{r}<=DataRef+90,"Até 90 dias",'
              f'IF(O{r}<=EDATE(DataRef,12),"91 dias a 12 meses","Após 12 meses")))))'),
        'T': f'=IF(O{r}="","",DATE(YEAR(O{r}),MONTH(O{r}),1))',
        'U': (f'=IF(R{r}="Ação imediata",1,IF(N{r}="Não encontrado",2,IF(N{r}="Revisão",3,IF(N{r}="Em Elaboração",4,'
              f'IF(S{r}="Até 90 dias",5,IF(R{r}="Validar cadastro",6,0))))))'),
        'V': (f'=CHOOSE(U{r}+1,IF(R{r}="Planejar revisão","Incluir no plano anual de revisão",IF(R{r}="Inativo","Sem ação","Monitorar")),'
              f'"Revisar ou prorrogar validade","Confirmar status no NDocs",'
              f'"Acompanhar revisão","Concluir elaboração",'
              f'"Iniciar revisão (≤ 90 dias)","Validar aplicabilidade")'),
        'W': (f'=(OR(fArea="Todas",ISNUMBER(SEARCH(fArea,G{r}))))*(OR(fStatus="Todos",N{r}=fStatus))*(OR(fSit="Todas",R{r}=fSit))'
              f'*(OR(fResp="Todos",P{r}=fResp))*(OR(fHor="Todos",S{r}=fHor))'),
        'X': f'=IF(AND(W{r}=1,U{r}>0),(7-U{r})*100000+50000-IF(ISNUMBER(Q{r}),Q{r},0)+ROW()/1000,0)',
        'Y': f'=UPPER(A{r}&" "&{s(f"C{r}")}&" "&{s(f"D{r}")})',
        'Z': f'=IF(qBusca="","",IF(ISNUMBER(SEARCH(TRIM(qBusca),Y{r})),ROW(),""))',
    }
    for cl, fx in f.items(): bs[f'{cl}{r}'] = fx
    for i in range(1, 27):
        c = bs.cell(r, i)
        c.font = F(9, color=MUTED if i >= 23 else TXT)
        c.border = Border(bottom=side())
        c.alignment = AL('center' if i in (2, 7, 13, 17, 21, 23) else 'left', 'center', i in (4,))
        if i >= 14: c.fill = FILL('F7F9FC')
    bs[f'F{r}'].number_format = 'dd/mm/yyyy'; bs[f'O{r}'].number_format = 'dd/mm/yyyy'
    bs[f'T{r}'].number_format = 'mmm/yyyy'; bs[f'Q{r}'].number_format = '#,##0;[Color10]-#,##0'
    bs[f'X{r}'].number_format = '0'
bs.freeze_panes = 'B5'
bs.auto_filter.ref = f'A4:Z{R1}'
bs.sheet_view.showGridLines = False
bs.sheet_view.zoomScale = 90
# realce discreto apenas da Situação Gerencial
for txt, fc, bc in [('Ação imediata', RED, REDL), ('Validar cadastro', RED, REDL), ('Em tratamento', AMB, AMBL),
                    ('Planejar revisão', AMB, AMBL), ('Inativo', GRY, GRYL), ('Regular', GRN, None)]:
    bs.conditional_formatting.add(f'R{R0}:R{R1}', CellIsRule(operator='equal', formula=[f'"{txt}"'],
                                  font=Font(color=fc, bold=txt != 'Regular'), fill=FILL(bc) if bc else None))
bs.conditional_formatting.add(f'A{R0}:A{R1}', FormulaRule(formula=[f'$M{R0}>1'], font=Font(color=BLUE, bold=True)))
nav_b = [('← Início', 'INÍCIO'), ('Dashboard', DB), ('Pesquisa', 'PESQUISA'), ('Qualidade', 'QUALIDADE BASE')]
for i, (t, s_) in enumerate(nav_b):
    c = bs.cell(1, 5 + i * 1 + (i > 0) * 0)
for i, (t, s_) in enumerate(nav_b):
    c = bs.cell(2 if False else 1, [6, 9, 12, 14][i]); link(c, s_, t); c.font = F(9, True, BLUE)

# ---------- AUX: KPI / agenda / área / responsáveis / ação ----------
def CS(**kw):  # COUNTIFS com filtro
    parts = [f"{col('W')},1"] + [f"{col(k)},{v}" for k, v in kw.items()]
    return 'COUNTIFS(' + ','.join(parts) + ')'

# Agenda 12 meses (linha 20+)
ax['A18'] = 'AGENDA — próximos 12 meses (somente Aprovadas, validade ≥ data de referência)'; ax['A18'].font = F(10, True, NAVY)
for j, h in enumerate(['Mês', 'Rótulo', 'Qtd.', 'Normal', 'Concentração', 'Limite']): ax.cell(19, 1 + j, h).font = F(9, True, MUTED)
for m in range(12):
    r = 20 + m
    ax[f'A{r}'] = f'=DATE(YEAR(DataRef),MONTH(DataRef)+{m},1)'; ax[f'A{r}'].number_format = 'mmm/yyyy'
    ax[f'B{r}'] = (f'=CHOOSE(MONTH(A{r}),"jan","fev","mar","abr","mai","jun","jul","ago","set","out","nov","dez")&"/"&RIGHT(YEAR(A{r}),2)')
    ax[f'C{r}'] = f'=COUNTIFS({col("W")},1,{col("T")},A{r},{col("N")},"Aprovado",{col("O")},">="&DataRef)'
    ax[f'E{r}'] = f'=IF(AND(C{r}>0,C{r}>=$F$20),C{r},0)'
    ax[f'D{r}'] = f'=C{r}-E{r}'
    ax[f'C{r}'].number_format = '0'
    for cl in 'DE': ax[f'{cl}{r}'].number_format = '0;;;'
ax['F20'] = '=MAX(3,ROUNDUP(2*AVERAGE(C20:C31),0))'
ax['G20'] = 'Mês de concentração = ≥ 2× a média mensal dos 12 meses e no mínimo 3 documentos'
ax['G20'].font = F(9, color=MUTED, italic=True)

# Área (linhas 35+)
area_order = [a for a, _ in collections.Counter(area_union(l) for l in groups.values()).most_common()]
ax['A34'] = 'CARTEIRA POR ÁREA CONSOLIDADA (cada código conta 1 vez; compartilhados em categoria própria, ex.: Q4/AOL)'
ax['A34'].font = F(10, True, NAVY)
for i, a in enumerate(area_order):
    ax.cell(35 + i, 1, a); ax.cell(35 + i, 2, f'=COUNTIFS({col("W")},1,{col("G")},A{35 + i})')
AREA_END = 35 + len(area_order) - 1
ax.cell(AREA_END + 1, 1, 'Total'); ax.cell(AREA_END + 1, 2, f'=SUM(B35:B{AREA_END})')

# Responsáveis (linhas 50+)
names = resps[1:]
ax['A48'] = 'CARTEIRA POR RESPONSÁVEL (nome padronizado — PROPER/TRIM); ordem = necessidade de acompanhamento'
ax['A48'].font = F(10, True, NAVY)
for j, h in enumerate(['Responsável', 'Total', 'Em tratamento', 'Ação/validar', 'Vencem 12m', 'Índice ordenação', '', 'k', 'Score', 'Linha']):
    ax.cell(49, 1 + j, h).font = F(9, True, MUTED)
RS = 50; RE = RS + len(names) - 1
for i, nm in enumerate(names):
    r = RS + i
    ax[f'A{r}'] = nm
    ax[f'B{r}'] = f'=COUNTIFS({col("W")},1,{col("P")},A{r})'
    ax[f'C{r}'] = f'=COUNTIFS({col("W")},1,{col("P")},A{r},{col("R")},"Em tratamento")'
    ax[f'D{r}'] = (f'=COUNTIFS({col("W")},1,{col("P")},A{r},{col("R")},"Ação imediata")'
                   f'+COUNTIFS({col("W")},1,{col("P")},A{r},{col("R")},"Validar cadastro")')
    ax[f'E{r}'] = f'=COUNTIFS({col("W")},1,{col("P")},A{r},{col("R")},"Planejar revisão")'
    ax[f'F{r}'] = f'=IF(B{r}=0,0,D{r}*1000000+C{r}*10000+E{r}*100+B{r}+{(len(names) - i)}/1000)'
for k in range(1, 9):
    r = RS + k - 1
    ax[f'H{r}'] = k
    ax[f'I{r}'] = f'=LARGE($F${RS}:$F${RE},H{r})'
    ax[f'J{r}'] = f'=IF(I{r}<=0,"",MATCH(I{r},$F${RS}:$F${RE},0))'

# Ação necessária (linhas 75+)
ax['A73'] = 'AÇÃO NECESSÁRIA — top 8 por prioridade (1 vencida ativa · 2 não encontrada · 3 revisão · 4 elaboração · 5 vence ≤ 90 dias · 6 validar cadastro)'
ax['A73'].font = F(10, True, NAVY)
for j, h in enumerate(['k', 'Score', 'Linha base']): ax.cell(74, 1 + j, h).font = F(9, True, MUTED)
for k in range(1, 9):
    r = 74 + k
    ax[f'A{r}'] = k
    ax[f'B{r}'] = f'=LARGE({col("X")},A{r})'
    ax[f'C{r}'] = f'=IF(B{r}<=0,"",MATCH(B{r},{col("X")},0))'
for c_, w_ in zip('ABCDEFGHIJ', [30, 12, 14, 14, 16, 18, 30, 6, 12, 8]): ax.column_dimensions[c_].width = w_
ax.sheet_view.showGridLines = False
ax.sheet_properties.tabColor = 'BFC6D1'

# ---------- DASHBOARD ----------
ds = wb.create_sheet(DB)
ds.sheet_view.showGridLines = False; ds.sheet_view.zoomScale = 100
ds.column_dimensions['A'].width = 2.5
for i in range(2, 26): ds.column_dimensions[L(i)].width = 7.9
ds.column_dimensions['Z'].width = 2.5
for r in range(1, 60):
    for c in range(1, 27): ds.cell(r, c).fill = FILL(BG)
heights = {1: 8, 2: 30, 3: 20, 4: 22, 5: 10, 6: 16, 7: 24, 8: 12, 9: 20, 10: 40, 11: 18, 12: 8, 13: 16,
           14: 22, 15: 16, 16: 22, 25: 16, 26: 14, 27: 22, 28: 16, 29: 20, 38: 16, 39: 14, 40: 22, 41: 18, 42: 26, 43: 18,
           44: 12, 45: 15, 46: 15, 47: 15, 48: 15}
for r, h in heights.items(): ds.row_dimensions[r].height = h
for r in range(17, 25): ds.row_dimensions[r].height = 30
for r in range(30, 38): ds.row_dimensions[r].height = 19

# cabeçalho executivo
box(ds, 'B2:Y3', fill=FILL(NAVY))
mput(ds, 'B2:N2', 'GESTÃO DOCUMENTAL — LCQ RJ', F(17, True, WHITE), align=AL('left', 'bottom', ind=1))
mput(ds, 'B3:N3', 'Cockpit Executivo de Instruções', F(10, False, 'C9D3E6'), align=AL('left', 'top', ind=1))
mput(ds, 'R2:U2', 'DATA DE REFERÊNCIA', F(8, True, '9FB0CC'), align=AL('right', 'bottom'))
c = mput(ds, 'R3:U3', '=DataRef', F(10, True, WHITE), align=AL('right', 'top')); c.number_format = 'dd/mm/yyyy'
mput(ds, 'V2:Y2', 'FONTE', F(8, True, '9FB0CC'), align=AL('right', 'bottom', ind=1))
mput(ds, 'V3:Y3', 'NDocs Rev.28', F(10, True, WHITE), align=AL('right', 'top', ind=1))
box(ds, 'B4:Y4', fill=FILL(WHITE), border=Border(bottom=side(LINE)))
navs = [('B4:D4', '⌂  Início', 'INÍCIO'), ('E4:H4', 'Base Consolidada', 'BASE CONSOLIDADA'), ('I4:K4', 'Pesquisa', 'PESQUISA'),
        ('L4:O4', 'Qualidade da Base', 'QUALIDADE BASE'), ('P4:S4', 'Outros Controles', 'OUTROS CONTROLES')]
for rng, t, s_ in navs:
    c = mput(ds, rng, None, F(9, True, BLUE), align=AL('left', 'center', ind=1)); link(c, s_, t); c.font = F(9, True, BLUE)
mput(ds, 'T4:Y4', 'Filtros afetam todos os blocos', F(8, color=MUTED, italic=True), align=AL('right', 'center', ind=1))

# filtros
filt = [('B6:E6', 'B7:E7', 'ÁREA', 'Todas', 'Áreas'), ('F6:I6', 'F7:I7', 'STATUS NDOCS', 'Todos', 'Status'),
        ('J6:M6', 'J7:M7', 'SITUAÇÃO GERENCIAL', 'Todas', 'Situações'), ('N6:R6', 'N7:R7', 'RESPONSÁVEL', 'Todos', 'Responsáveis'),
        ('S6:V6', 'S7:V7', 'HORIZONTE DE VALIDADE', 'Todos', 'Horizontes')]
for lab, inp, t, dflt, lst in filt:
    mput(ds, lab, t, F(8, True, MUTED), align=AL('left', 'bottom'))
    c = mput(ds, inp, dflt, F(10, True, NAVY), align=AL('left', 'center', ind=1))
    box(ds, inp, fill=FILL(WHITE), border=Border(top=side(LINE), bottom=side(BLUE2, 'medium'), left=side(LINE), right=side(LINE)))
    dv = DataValidation(type='list', formula1=f'={LR[lst]}', allow_blank=False, showDropDown=False)
    dv.error = 'Escolha um valor da lista.'; dv.errorTitle = 'Filtro'; dv.prompt = 'Clique na seta para filtrar'; dv.showErrorMessage = True
    ds.add_data_validation(dv); dv.add(inp.split(':')[0])
c = mput(ds, 'W6:Y7', '↺ Para limpar, volte\ncada filtro a "Todas/Todos"', F(8, color=MUTED, italic=True), align=AL('left', 'center', True))

# KPI cards
ACT = f'COUNTIFS({col("W")},1,{col("R")},"<>Inativo")'
cards = [
    ('B', 'INSTRUÇÕES ÚNICAS', f'=SUM({col("W")})', '0',
     f'={ACT}&" ativas · "&COUNTIFS({col("W")},1,{col("R")},"Inativo")&" inativas"', NAVY),
    ('F', 'APROVADAS', f'=COUNTIFS({col("W")},1,{col("N")},"Aprovado")', '0',
     f'=IFERROR(TEXT(COUNTIFS({col("W")},1,{col("N")},"Aprovado")/{ACT},"0%"),"–")&" das instruções ativas"', GRN),
    ('J', 'EM PROCESSO', f'=COUNTIFS({col("W")},1,{col("R")},"Em tratamento")', '0',
     f'="Revisão "&COUNTIFS({col("W")},1,{col("N")},"Revisão")&" · Elaboração "&COUNTIFS({col("W")},1,{col("N")},"Em Elaboração")', AMB),
    ('N', 'AÇÃO IMEDIATA', f'=COUNTIFS({col("W")},1,{col("R")},"Ação imediata")', '0',
     f'="vencida(s) e ativa(s) · +"&COUNTIFS({col("W")},1,{col("R")},"Validar cadastro")&" a validar"', RED),
    ('R', 'VENCEM EM 12 MESES', f'=COUNTIFS({col("W")},1,{col("N")},"Aprovado",{col("O")},">="&DataRef,{col("O")},"<="&EDATE(DataRef,12))', '0',
     f'="até 90 dias: "&COUNTIFS({col("W")},1,{col("N")},"Aprovado",{col("S")},"Até 90 dias")', BLUE),
    ('V', 'SAÚDE DOCUMENTAL', f'=IFERROR((COUNTIFS({col("W")},1,{col("R")},"Regular")+COUNTIFS({col("W")},1,{col("R")},"Planejar revisão"))/{ACT},"–")', '0.0%',
     '="aprovadas e válidas ÷ ativas"', NAVY),
]
for c0, lab, val, fmt, sub, acc in cards:
    i0 = openpyxl.utils.column_index_from_string(c0); c3 = L(i0 + 3)
    rng = f'{c0}9:{c3}12'
    for r in range(9, 13):
        for ci in range(i0, i0 + 4):
            cell = ds.cell(r, ci); cell.fill = FILL(WHITE)
            cell.border = Border(left=side(BG, 'thick') if ci == i0 else None, right=side(BG, 'thick') if ci == i0 + 3 else None,
                                 top=side(acc, 'thick') if r == 9 else None)
    mput(ds, f'{c0}9:{c3}9', lab, F(8, True, MUTED), align=AL('left', 'bottom', ind=1))
    c = mput(ds, f'{c0}10:{c3}10', val, F(24, True, acc if acc != BLUE else NAVY), align=AL('left', 'center', ind=1)); c.number_format = fmt
    mput(ds, f'{c0}11:{c3}11', sub, F(8, color=MUTED), align=AL('left', 'top', ind=1))
# KPI ação imediata: vermelho só se > 0
ds.conditional_formatting.add('N10', CellIsRule(operator='equal', formula=['0'], font=Font(color=GRN)))
ds.conditional_formatting.add('V10', CellIsRule(operator='lessThan', formula=['0.9'], font=Font(color=AMB)))

def section(rng, title, cap_rng=None, cap=None):
    mput(ds, rng, title, F(11, True, NAVY), align=AL('left', 'bottom'))
    for row in ds[rng]:
        for c in row: c.border = Border(bottom=side(NAVY, 'thin'))
    if cap: mput(ds, cap_rng, cap, F(8, color=MUTED, italic=True), align=AL('left', 'center'))

# Ação necessária agora
section('B14:P14', 'AÇÃO NECESSÁRIA AGORA', 'B15:P15',
        'Somente o que exige tratamento. Ordem: 1 vencida ativa · 2 não encontrada · 3 revisão · 4 elaboração · 5 vence ≤ 90 dias · 6 validar cadastro')
cols_a = [('B16:B16', 'P', 'U'), ('C16:E16', 'Código', 'A'), ('F16:I16', 'Título', 'D'), ('J16:J16', 'Área', 'G'),
          ('K16:L16', 'Situação', 'R'), ('M16:N16', 'Responsável', 'P'), ('O16:P16', 'Ação recomendada', 'V')]
for rng, h, src in cols_a:
    c = mput(ds, rng, h, F(8, True, MUTED), align=AL('center' if h in ('P', 'Área') else 'left', 'center', ind=0 if h in ('P', 'Área') else 1))
    box(ds, rng, fill=FILL(WHITE), border=Border(bottom=side(LINE)))
for k in range(8):
    r = 17 + k; ar = 75 + k
    for rng, h, src in cols_a:
        a, b = rng.replace('16', str(r)).split(':')
        fx = f'=IF({AX}!$C${ar}="","",INDEX({col(src)},{AX}!$C${ar}))'
        if k == 0 and h == 'Código':
            fx = f'=IF({AX}!$C${ar}="","✓ Nenhum item exige ação com os filtros atuais",INDEX({col(src)},{AX}!$C${ar}))'
        c = mput(ds, f'{a}:{b}', fx, F(9 if h != 'Título' else 8, h in ('P', 'Código'), TXT),
                 align=AL('center' if h in ('P', 'Área') else 'left', 'center', h in ('Título', 'Ação recomendada', 'Responsável'), 0 if h in ('P', 'Área') else 1))
        box(ds, f'{a}:{b}', fill=FILL(WHITE), border=Border(bottom=side(LINE)))
        if h == 'P':
            ds.conditional_formatting.add(a, CellIsRule(operator='lessThanOrEqual', formula=['2'], font=Font(color=WHITE, bold=True), fill=FILL(RED)))
            ds.conditional_formatting.add(a, CellIsRule(operator='between', formula=['3', '5'], font=Font(color=WHITE, bold=True), fill=FILL('C98A12')))
            ds.conditional_formatting.add(a, CellIsRule(operator='equal', formula=['6'], font=Font(color=WHITE, bold=True), fill=FILL(GRY)))
        if h == 'Situação':
            for txt, fc in [('Ação imediata', RED), ('Validar cadastro', RED), ('Em tratamento', AMB), ('Planejar revisão', AMB)]:
                ds.conditional_formatting.add(a, CellIsRule(operator='equal', formula=[f'"{txt}"'], font=Font(color=fc, bold=True)))

# Agenda (gráfico)
section('R14:Y14', 'AGENDA DE VENCIMENTOS — 12 MESES', 'R15:Y15', 'Aprovadas que vencem por mês · barra escura = mês de concentração')
box(ds, 'R16:Y25', fill=FILL(WHITE))
ch = BarChart(); ch.type = 'col'; ch.grouping = 'stacked'; ch.overlap = 100; ch.gapWidth = 45
data = Reference(ax, min_col=4, max_col=5, min_row=19, max_row=31)
ch.add_data(data, titles_from_data=True)
ch.set_categories(Reference(ax, min_col=2, min_row=20, max_row=31))
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties
for s_, colr, lc in zip(ch.series, [BLUE2, NAVY], [NAVY, WHITE]):
    s_.graphicalProperties = GraphicalProperties(solidFill=colr); s_.graphicalProperties.line.noFill = True
    s_.dLbls = DataLabelList(); s_.dLbls.showVal = True
    for a_ in ('showSerName', 'showCatName', 'showLegendKey', 'showPercent'): setattr(s_.dLbls, a_, False)
    s_.dLbls.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=900, b=True, solidFill=lc)), endParaRPr=CharacterProperties())])
ch.legend.position = 'b'
ch.y_axis.delete = True; ch.y_axis.majorGridlines = None
ch.x_axis.delete = False; ch.x_axis.tickLblPos = 'low'
ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=LINE))
ch.width = 14.2; ch.height = 6.9
ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
ds.add_chart(ch, 'R16')

# Carteira por área (gráfico)
section('B27:H27', 'CARTEIRA POR ÁREA', 'B28:H28', 'Código único; compartilhados em categoria própria (sem dupla contagem)')
box(ds, 'B29:H38', fill=FILL(WHITE))
ca = BarChart(); ca.type = 'bar'; ca.gapWidth = 40
ca.add_data(Reference(ax, min_col=2, min_row=35, max_row=AREA_END), titles_from_data=False)
ca.set_categories(Reference(ax, min_col=1, min_row=35, max_row=AREA_END))
s_ = ca.series[0]; s_.graphicalProperties = GraphicalProperties(solidFill=BLUE); s_.graphicalProperties.line.noFill = True
s_.dLbls = DataLabelList(); s_.dLbls.showVal = True
for a_ in ('showSerName', 'showCatName', 'showLegendKey', 'showPercent'): setattr(s_.dLbls, a_, False)
ca.legend = None; ca.y_axis.delete = True; ca.y_axis.majorGridlines = None
ca.x_axis.delete = False; ca.x_axis.scaling.orientation = 'maxMin'
ca.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=LINE))
ca.width = 11.6; ca.height = 5.6
ds.add_chart(ca, 'B29')

# Matriz Área × Situação
section('J27:Q27', 'MATRIZ ÁREA × SITUAÇÃO', 'J28:Q28', 'Concentração de risco por área consolidada')
msits = ['Ação imediata', 'Validar cadastro', 'Em tratamento', 'Planejar revisão', 'Regular', 'Inativo']
mhead = ['Ação imed.', 'Validar', 'Em trat.', 'Planejar', 'Regular', 'Inativo']
c = mput(ds, 'J29:K29', 'Área', F(8, True, MUTED), align=AL('left', 'center', ind=1))
for j, h in enumerate(mhead):
    c = ds.cell(29, 12 + j, h); c.font = F(8, True, MUTED); c.alignment = AL('center', 'center', True)
box(ds, 'J29:Q29', fill=FILL(WHITE), border=Border(bottom=side(LINE)))
for i, a in enumerate(area_order + ['Total']):
    r = 30 + i
    c = mput(ds, f'J{r}:K{r}', a, F(9, a == 'Total', TXT), align=AL('left', 'center', ind=1))
    for j, sname in enumerate(msits):
        cc = ds.cell(r, 12 + j)
        cc.value = (f'=COUNTIFS({col("W")},1,{col("G")},$J{r},{col("R")},"{sname}")' if a != 'Total'
                    else f'=SUM({L(12 + j)}30:{L(12 + j)}{r - 1})')
        cc.number_format = '0;-0;"·"'; cc.font = F(9, a == 'Total', TXT); cc.alignment = AL('center')
    box(ds, f'J{r}:Q{r}', fill=FILL(WHITE), border=Border(bottom=side(LINE if a != 'Total' else MUTED)))
MR = 30 + len(area_order) - 1
ds.conditional_formatting.add(f'L30:M{MR}', CellIsRule(operator='greaterThan', formula=['0'], fill=FILL(REDL), font=Font(color=RED, bold=True)))
ds.conditional_formatting.add(f'N30:O{MR}', CellIsRule(operator='greaterThan', formula=['0'], fill=FILL(AMBL), font=Font(color=AMB, bold=True)))
ds.conditional_formatting.add(f'P30:P{MR}', ColorScaleRule(start_type='num', start_value=0, start_color=WHITE, end_type='max', end_color='C9D6EC'))
ds.conditional_formatting.add(f'Q30:Q{MR}', CellIsRule(operator='greaterThan', formula=['0'], font=Font(color=GRY)))

# Carteira por responsável
section('S27:Y27', 'CARTEIRA POR RESPONSÁVEL', 'S28:Y28', 'Distribuição para acompanhamento — não é ranking de desempenho')
for rng, h in [('S29:U29', 'Responsável'), ('V29:V29', 'Total'), ('W29:W29', 'Em trat.'), ('X29:X29', 'Ação/ valid.'), ('Y29:Y29', '12 m')]:
    c = mput(ds, rng, h, F(8, True, MUTED), align=AL('left' if h == 'Responsável' else 'center', 'center', True, 1 if h == 'Responsável' else 0))
    box(ds, rng, fill=FILL(WHITE), border=Border(bottom=side(LINE)))
for k in range(8):
    r = 30 + k; ar = RS + k
    idx = f'{AX}!$J${ar}'
    c = mput(ds, f'S{r}:U{r}', f'=IF({idx}="","",INDEX({AX}!$A${RS}:$A${RE},{idx}))', F(8, False, TXT), align=Alignment(horizontal='left', vertical='center', shrink_to_fit=True, indent=1))
    for ci, sc in zip('VWXY', 'BCDE'):
        cc = ds[f'{ci}{r}']; cc.value = f'=IF({idx}="","",INDEX({AX}!${sc}${RS}:${sc}${RE},{idx}))'
        cc.font = F(9, ci == 'V', TXT); cc.alignment = AL('center'); cc.number_format = '0;-0;"·"'
    box(ds, f'S{r}:Y{r}', fill=FILL(WHITE), border=Border(bottom=side(LINE)))
ds.conditional_formatting.add('X30:X37', CellIsRule(operator='greaterThan', formula=['0'], fill=FILL(REDL), font=Font(color=RED, bold=True)))
ds.conditional_formatting.add('W30:W37', CellIsRule(operator='greaterThan', formula=['0'], fill=FILL(AMBL), font=Font(color=AMB, bold=True)))
mput(ds, 'S38:Y38', f'=COUNTIF({AX}!$B${RS}:$B${RE},">0")&" responsáveis com carteira · lista completa em _AUX DASH"',
     F(8, color=MUTED, italic=True), align=AL('left', 'center'))

# Outros controles (faixa)
section('B40:Y40', 'OUTROS CONTROLES DOCUMENTAIS  ·  não somados aos KPIs de instruções')
fams = [('NDocs FMG e ANX', 'B', 3), ('Documentos Transversais', 'B', 3), ('LPP', 'A', 3), ('AST AOL', 'A', 3),
        ('AST Geral', 'A', 3), ("RT's", 'B', 3), ('MO', 'B', 3)]
def famref(n, c_, r0):
    ws_ = wb[n]
    q = n.replace("'", "''")
    return f"'{q}'!${c_}${r0}:${c_}${ws_.max_row}", f"'{q}'!$A${r0}:${L(min(ws_.max_column, 14))}${ws_.max_row}"
for i, (n, c_, r0) in enumerate(fams):
    a = 2 + i * 3 + (1 if i >= 0 else 0) * 0
    a = 2 + round(i * 24 / 7); b = 2 + round((i + 1) * 24 / 7) - 1
    A_, B_ = L(a), L(b)
    kref, eref = famref(n, c_, r0)
    for r in range(41, 44):
        for ci in range(a, b + 1):
            ds.cell(r, ci).fill = FILL(WHITE)
            ds.cell(r, ci).border = Border(left=side(BG, 'thick') if ci == a else None, right=side(BG, 'thick') if ci == b else None)
    c = mput(ds, f'{A_}41:{B_}41', n, F(8, True, MUTED), align=AL('left', 'bottom', ind=1))
    c = mput(ds, f'{A_}42:{B_}42', f'=COUNTA({kref})', F(15, True, NAVY), align=AL('left', 'center', ind=1)); c.number_format = '0'
    c = mput(ds, f'{A_}43:{B_}43', f'=IF(SUMPRODUCT(--ISERROR({eref}))>0,"⚠ "&SUMPRODUCT(--ISERROR({eref}))&" erro(s) legado(s)","Abrir aba →")',
             F(8, True, BLUE), align=AL('left', 'top', ind=1))
    link(c, n, None); c.hyperlink.display = n
    ds.conditional_formatting.add(f'{A_}43', FormulaRule(formula=[f'LEFT({A_}43,1)="⚠"'], font=Font(color=AMB)))

# Notas
notes = [
    'Saúde documental = (Regular + Planejar revisão) ÷ instruções ativas. Ativas = todas, exceto Canceladas e Fora de Uso. Mede "quantas instruções em uso estão aprovadas e dentro da validade".',
    'Situação Gerencial é camada analítica: o Status oficial do NDocs é preservado na Base Consolidada (coluna E). Regras completas em QUALIDADE BASE › Regras.',
    'Contagem por Código NDocs único (125 códigos de 131 linhas de instrução). Documento compartilhado entre áreas conta uma vez; o filtro de Área inclui os compartilhados (ex.: Q4 inclui Q4/AOL).',
]
for i, t in enumerate(notes):
    mput(ds, f'B{45 + i}:Y{45 + i}', t, F(8, color=MUTED), align=AL('left', 'center'))
ds.print_area = 'A1:Z48'
ds.page_setup.orientation = 'landscape'; ds.page_setup.fitToWidth = 1; ds.page_setup.fitToHeight = 1
ds.sheet_properties.pageSetUpPr.fitToPage = True
ds.sheet_properties.tabColor = NAVY

# ---------- PESQUISA ----------
ps = wb.create_sheet('PESQUISA')
ps.sheet_view.showGridLines = False
ps.column_dimensions['A'].width = 2.5
for c_, w_ in zip('BCDEFGHIJKL', [5, 22, 16, 50, 8, 16, 13, 9, 30, 18, 2.5]): ps.column_dimensions[c_].width = w_
for r in range(1, 30):
    for ci in range(1, 13): ps.cell(r, ci).fill = FILL(BG)
box(ps, 'B2:K3', fill=FILL(NAVY))
mput(ps, 'B2:H2', 'PESQUISA DE DOCUMENTO', F(15, True, WHITE), align=AL('left', 'bottom', ind=1))
mput(ps, 'B3:H3', 'Busca por Código NDocs, número antigo (IT …) ou parte do título — sem diferenciar maiúsculas', F(9, color='C9D3E6'), align=AL('left', 'top', ind=1))
c = mput(ps, 'I2:K3', None, F(9, True, WHITE), align=AL('right', 'center', ind=1)); link(c, DB, '← Dashboard'); c.font = F(9, True, WHITE)
ps.row_dimensions[2].height = 28; ps.row_dimensions[3].height = 20
mput(ps, 'B5:E5', 'DIGITE O TERMO DE BUSCA', F(8, True, MUTED), align=AL('left', 'bottom'))
ps['B6'] = '🔍'; ps['B6'].alignment = AL('center'); ps['B6'].fill = FILL(WHITE)
ps['C6'].fill = FILL(WHITE)
mput(ps, 'C6:C6', 'Termo:', F(9, True, MUTED), align=AL('right'))
c = mput(ps, 'D6:E6', 'OXIGÊNIO', F(12, True, NAVY), align=AL('left', 'center', ind=1))
box(ps, 'B6:E6', fill=FILL(WHITE), border=Border(bottom=side(BLUE, 'medium'), top=side(LINE)))
ps.row_dimensions[6].height = 28
mput(ps, 'F6:K6', f'=IF(qBusca="","Digite um termo para pesquisar",COUNT({col("Z")})&" resultado(s) · exibindo até 15")', F(9, True, BLUE), align=AL('left', 'center', ind=1))
ps.row_dimensions[8].height = 28
hp = ['#', 'Código NDocs', 'Antigo nº', 'Título', 'Versão', 'Status NDocs', 'Validade', 'Área', 'Responsável', 'Situação Gerencial']
srcp = [None, 'A', 'C', 'D', 'B', 'E', 'O', 'G', 'P', 'R']
for j, h in enumerate(hp):
    c = ps.cell(8, 2 + j, h); c.font = F(8, True, WHITE); c.fill = FILL(BLUE); c.alignment = AL('left', 'center', True, 1)
for k in range(1, 16):
    r = 8 + k
    ps.row_dimensions[r].height = 30
    ps[f'B{r}'] = f'=IF(COUNT({col("Z")})<{k},"",{k})'
    ps[f'B{r}'].alignment = AL('center')
    for j, sc in enumerate(srcp[1:], 1):
        cc = ps.cell(r, 2 + j)
        cc.value = f'=IF($B{r}="","",INDEX({BS}!${sc}$1:${sc}${R1},SMALL({col("Z")},$B{r})))'
        cc.alignment = AL('left', 'center', sc in ('D', 'P'), 1)
    for j in range(10):
        cc = ps.cell(r, 2 + j); cc.fill = FILL(WHITE); cc.border = Border(bottom=side(LINE)); cc.font = F(9, j == 1, TXT)
    ps[f'H{r}'].number_format = 'dd/mm/yyyy'
for txt, fc in [('Ação imediata', RED), ('Validar cadastro', RED), ('Em tratamento', AMB), ('Planejar revisão', AMB), ('Inativo', GRY), ('Regular', GRN)]:
    ps.conditional_formatting.add('K9:K23', CellIsRule(operator='equal', formula=[f'"{txt}"'], font=Font(color=fc, bold=True)))
mput(ps, 'B25:K25', 'Dica: digite só o número (ex.: 016238 ou 5020-00217) ou uma palavra do título. Apague o termo para limpar. Detalhes completos na Base Consolidada.',
     F(8, color=MUTED, italic=True), align=AL('left'))
ps.freeze_panes = 'A9'
ps.sheet_properties.tabColor = NAVY

# ---------- QUALIDADE BASE ----------
qs = wb.create_sheet('QUALIDADE BASE')
qs.sheet_view.showGridLines = False
qs.column_dimensions['A'].width = 2.5
for c_, w_ in zip('BCDEFGH', [7, 13, 26, 26, 60, 30, 46]): qs.column_dimensions[c_].width = w_
box(qs, 'B2:H3', fill=FILL(NAVY))
mput(qs, 'B2:F2', 'QUALIDADE E CONFIABILIDADE DA BASE', F(15, True, WHITE), align=AL('left', 'bottom', ind=1))
mput(qs, 'B3:F3', 'Auditoria da consolidação das Instruções — achados classificados em Crítico, Atenção e Informativo', F(9, color='C9D3E6'), align=AL('left', 'top', ind=1))
c = mput(qs, 'G2:H3', None, F(9, True, WHITE), align=AL('right', 'center', ind=1)); link(c, DB, '← Dashboard'); c.font = F(9, True, WHITE)
qs.row_dimensions[2].height = 28; qs.row_dimensions[3].height = 20

# achados (snapshot gerado a partir das origens)
fd = []
def add(cls, cat, ref, desc, ev, trat): fd.append((cls, cat, ref, desc, ev, trat))
for code, l in groups.items():
    if len(l) > 1:
        sheets = {x['sheet'] for x in l}
        orig = '; '.join(f"{x['sheet']}!{x['row']}" for x in l)
        if len(sheets) > 1:
            add('Informativo', 'Duplicidade esperada', code, f'Instrução compartilhada entre áreas (área consolidada {area_union(l)})', orig, 'Contada uma vez; área consolidada preserva o compartilhamento')
        else:
            add('Atenção', 'Duplicidade suspeita', code, 'Mesmo código repetido na mesma aba, sem justificativa aparente', orig, 'Contado uma vez; recomenda-se remover a linha repetida na origem')
        tits = {str(x['tit']).replace('\xa0', ' ').strip() for x in l}
        if len(tits) > 1:
            add('Atenção', 'Divergência cadastral — título', code, ' ≠ '.join(sorted(tits)), orig, 'Base usa o título da 1ª origem; alinhar com o NDocs')
        vals = {pdate(x['val']) for x in l}
        if len(vals) > 1: add('Atenção', 'Divergência cadastral — validade', code, 'Validades diferentes', orig, 'Base usa a 1ª origem')
        if len({x['ver'] for x in l}) > 1: add('Atenção', 'Divergência cadastral — versão', code, 'Versões diferentes', orig, 'Base usa a 1ª origem')
        if len({nstatus(x['st']) for x in l}) > 1: add('Atenção', 'Divergência cadastral — status', code, 'Status diferentes', orig, 'Base usa a 1ª origem')
        rs = {str(x['resp']).replace('\xa0', ' ').strip() for x in l}
        if len(rs) > 1 and len({r_.title() for r_ in rs}) == 1:
            add('Informativo', 'Divergência cadastral — grafia do responsável', code, ' / '.join(sorted(rs)), orig, 'Unificado na coluna "Responsável padronizado"')
for d in recs:
    if isinstance(d['val'], str):
        add('Atenção', 'Validade gravada como texto', d['cod'], f'Valor "{d["val"].strip()}" está como texto, não como data', f"{d['sheet']}!G{d['row']}", 'Convertida na coluna "Validade (analítica)" (dd/mm/aaaa); origem preservada')
    if d['val'] is None and nstatus(d['st']) != 'Em Elaboração':
        add('Crítico', 'Campo crítico vazio — validade', d['cod'], 'Instrução ativa sem validade', f"{d['sheet']}!G{d['row']}", 'Classificada como "Validar cadastro"')
    if d['val'] is None and nstatus(d['st']) == 'Em Elaboração':
        add('Informativo', 'Validade vazia (esperado)', d['cod'], 'Documento em elaboração ainda sem validade', f"{d['sheet']}!G{d['row']}", 'Sem impacto: situação "Em tratamento"')
    s0 = str(d['st'])
    if nstatus(d['st']) == 'Não encontrado':
        add('Crítico', 'Status crítico', d['cod'], 'Status "Não encontrado" — documento não localizado no NDocs', f"{d['sheet']}!F{d['row']}", 'Prioridade 2 no painel Ação Necessária')
    if s0.strip() == 'Avaiação de uso':
        add('Atenção', 'Status não padronizado', d['cod'], 'Grafia "Avaiação de uso" (status fora da lista NDocs); comentário: ' + str(d['com']).strip(), f"{d['sheet']}!F{d['row']}", 'Padronizado como "Avaliação de uso" → Validar cadastro')
    if str(d['tit']).startswith('ONTROLE'):
        add('Informativo', 'Possível erro de digitação', d['cod'], f'Título inicia com "{str(d["tit"])[:20]}…" (provável "CONTROLE")', f"{d['sheet']}!E{d['row']}", 'Não alterado; corrigir na origem')
sp_ = [d for d in recs if isinstance(d['st'], str) and d['st'] != d['st'].strip()]
add('Informativo', 'Status não padronizado', f'{len(sp_)} registros', 'Status com espaço no final (ex.: "Aprovado ")', ', '.join(sorted({d['sheet'] for d in sp_})), 'Removido na coluna "Status padronizado"')
cc_ = [d for d in recs if str(d['st']).strip() == 'Cancelada']
if cc_: add('Informativo', 'Status não padronizado', f'{len(cc_)} registro(s)', '"Cancelada" e "Cancelado" coexistem', f"{cc_[0]['sheet']}!F{cc_[0]['row']}", 'Unificado como "Cancelado"')
nb = [d for d in recs if any('\xa0' in str(d[k]) for k in ('ant', 'tit', 'val'))]
add('Informativo', 'Caractere invisível', f'{len(nb)} registros', 'Espaço não separável (NBSP) em Antigo nº / Título / Validade', ', '.join(sorted({d['sheet'] for d in nb})), 'Tratado na busca e na validade analítica')
rv = collections.defaultdict(set)
for d in recs: rv[str(d['resp']).replace('\xa0', ' ').strip().title()].add(str(d['resp']).replace('\xa0', ' ').strip())
for k, v in rv.items():
    if len(v) > 1: add('Informativo', 'Normalização de responsável', k, ' / '.join(sorted(v)), 'Abas de instrução', 'Unificado na camada analítica (PROPER/TRIM); valor original preservado')
for d in lvs:
    add('Informativo', 'Registro fora do escopo', d['cod'], f'Tipo "{d["tipo"]}" em aba de Instruções', f"{d['sheet']}!A{d['row']}", 'Excluído da contagem de Instruções')
add('Atenção', 'Erro legado de fórmula', 'LPP!F100:F124', '25 fórmulas UPPER() com referência quebrada (erro REF) — origem da referência perdida', 'LPP', 'Não corrigido (sem fonte segura); fora dos KPIs de Instruções')
add('Atenção', 'Erro legado de fórmula', 'AST Geral!E42', 'Fórmula UPPER() com referência quebrada (erro REF)', 'AST Geral', 'Não corrigido (sem fonte segura); fora dos KPIs de Instruções')
add('Informativo', 'Erro legado não reproduzido', 'MO!H8:H10', 'Versão anterior registrava erro VALUE; valores atuais retornam " " (validade vazia)', 'MO (Tabela11)', 'Sem erro no arquivo recebido; fórmula depende de HOJE()')
add('Informativo', 'Abas históricas', 'NDocs, NDocs (2), NDocs Salvação', 'Cópias de revisões anteriores (ex.: versões 6.0 e validades 2024)', 'Abas históricas', 'Preservadas; não usadas nos KPIs')
order = {'Crítico': 0, 'Atenção': 1, 'Informativo': 2}
fd.sort(key=lambda x: order[x[0]])

# KPI da qualidade (ao vivo)
QK = [('B5:C5', 'B6:C6', 'LINHAS LIDAS', len(recs), 'Q4 + PE_PP + AOL + COMUNS (sem LVS)'),
      ('D5:D5', 'D6:D6', 'CÓDIGOS ÚNICOS', f'=COUNTA({col("A")})', 'contagem principal'),
      ('E5:E5', 'E6:E6', 'CRÍTICO', '=COUNTIF($C$27:$C$200,"Crítico")', 'exige correção no NDocs/origem'),
      ('F5:F5', 'F6:F6', 'ATENÇÃO', '=COUNTIF($C$27:$C$200,"Atenção")', 'corrigir na próxima atualização'),
      ('G5:H5', 'G6:H6', 'INFORMATIVO', '=COUNTIF($C$27:$C$200,"Informativo")', 'tratado na camada analítica')]
for lab, val, t, v, sub in QK:
    mput(qs, lab, t, F(8, True, MUTED), align=AL('left', 'bottom', ind=1))
    c = mput(qs, val, v, F(20, True, NAVY), align=AL('left', 'center', ind=1))
    sr = lab.replace('5', '7')
    mput(qs, sr, sub, F(8, color=MUTED), align=AL('left', 'top', ind=1))
    for rr in (lab, val, sr): box(qs, rr, fill=FILL(BLUEL))
qs['E6'].font = F(20, True, RED); qs['F6'].font = F(20, True, AMB); qs['G6'].font = F(20, True, MUTED)
qs.row_dimensions[6].height = 34

# verificações automáticas
mput(qs, 'B9:H9', 'VERIFICAÇÕES AUTOMÁTICAS (recalculam com a base)', F(11, True, NAVY), align=AL('left', 'bottom'))
for c in qs['B9:H9'][0]: c.border = Border(bottom=side(NAVY))
checks = [
    ('Status com grafia diferente do padrão', f'=SUMPRODUCT(--({col("E")}<>{col("N")}))', 'Informativo'),
    ('Validade gravada como texto', f'=SUMPRODUCT(--ISTEXT({col("F")}))', 'Atenção'),
    ('Ativas (exceto em elaboração) sem validade', f'=SUMPRODUCT(({col("O")}="")*({col("R")}<>"Inativo")*({col("N")}<>"Em Elaboração"))', 'Crítico'),
    ('Responsável vazio', f'=COUNTBLANK({col("I")})', 'Crítico'),
    ('Status fora da lista (Não encontrado / Avaliação de uso)', f'=COUNTIF({col("N")},"Não encontrado")+COUNTIF({col("N")},"Avaliação de uso")', 'Crítico'),
    ('Códigos com mais de um registro na origem', f'=COUNTIF({col("M")},">1")', 'Informativo'),
    ('Erros de fórmula na camada analítica da base', f"=SUMPRODUCT(--ISERROR({BS}!$N${R0}:$Z${R1}))", 'Crítico'),
    ('Erros legados — LPP / AST Geral / MO', "=SUMPRODUCT(--ISERROR(LPP!$A$1:$N$124))+SUMPRODUCT(--ISERROR('AST Geral'!$A$1:$N$43))+SUMPRODUCT(--ISERROR(MO!$A$1:$N$20))", 'Atenção'),
]
for j, h in enumerate(['#', 'Resultado', 'Verificação', '', 'Classe se > 0']):
    c = qs.cell(10, 2 + j, h if h else None); c.font = F(8, True, MUTED)
for i, (t, fx, cls) in enumerate(checks):
    r = 11 + i
    qs[f'B{r}'] = i + 1; qs[f'C{r}'] = fx; mput(qs, f'D{r}:E{r}', t, F(9), align=AL('left'))
    qs[f'F{r}'] = f'=IF(C{r}=0,"OK",{chr(34)}{cls}{chr(34)})'
    for c in qs[f'B{r}:H{r}'][0]:
        c.border = Border(bottom=side())
        if c.coordinate[0] not in 'DE': c.font = F(9, c.coordinate[0] == 'C')
    qs[f'C{r}'].alignment = AL('center')
CE = 11 + len(checks) - 1
for txt, fc, bc in [('Crítico', RED, REDL), ('Atenção', AMB, AMBL), ('Informativo', MUTED, GRYL), ('OK', GRN, GRNL)]:
    qs.conditional_formatting.add(f'F11:F{CE}', CellIsRule(operator='equal', formula=[f'"{txt}"'], font=Font(color=fc, bold=True), fill=FILL(bc)))

# tabela de achados
mput(qs, 'B24:H24', f'ACHADOS DA AUDITORIA DE ORIGEM — {len(fd)} itens (gerado em {dt.date.today():%d/%m/%Y} a partir das abas de origem)', F(11, True, NAVY), align=AL('left', 'bottom'))
for c in qs['B24:H24'][0]: c.border = Border(bottom=side(NAVY))
mput(qs, 'B25:H25', 'Crítico = afeta a decisão ou a contagem · Atenção = inconsistência a corrigir na origem · Informativo = esperado ou já tratado na camada analítica',
     F(8, color=MUTED, italic=True), align=AL('left'))
for j, h in enumerate(['ID', 'Classe', 'Categoria', 'Código / local', 'Descrição', 'Evidência (origem)', 'Tratamento no cockpit']):
    c = qs.cell(26, 2 + j, h); c.font = F(8, True, WHITE); c.fill = FILL(BLUE); c.alignment = AL('left', 'center', True, 1)
for i, row in enumerate(fd):
    r = 27 + i
    vals = [f'Q{i + 1:02d}'] + list(row)
    for j, v in enumerate(vals):
        c = qs.cell(r, 2 + j, v); c.font = F(9, j == 1); c.alignment = AL('left', 'center', True, 1); c.border = Border(bottom=side())
    qs.row_dimensions[r].height = 30
QE = 27 + len(fd) - 1
for txt, fc, bc in [('Crítico', RED, REDL), ('Atenção', AMB, AMBL), ('Informativo', MUTED, GRYL)]:
    qs.conditional_formatting.add(f'C27:C{QE}', CellIsRule(operator='equal', formula=[f'"{txt}"'], font=Font(color=fc, bold=True), fill=FILL(bc)))
qs.auto_filter.ref = f'B26:H{QE}'

# regras
RR = QE + 3
mput(qs, f'B{RR}:H{RR}', 'REGRAS — SITUAÇÃO GERENCIAL, PRIORIDADE E SAÚDE DOCUMENTAL', F(11, True, NAVY), align=AL('left', 'bottom'))
for c in qs[f'B{RR}:H{RR}'][0]: c.border = Border(bottom=side(NAVY))
rules = [
    ('Situação', 'Inativo', 'Status padronizado = Cancelado ou Fora de Uso. Não penaliza a saúde documental.'),
    ('Situação', 'Em tratamento', 'Status = Revisão ou Em Elaboração.'),
    ('Situação', 'Validar cadastro', 'Status diferente de Aprovado (ex.: Não encontrado, Avaliação de uso) ou Aprovado sem validade.'),
    ('Situação', 'Ação imediata', 'Aprovado com validade anterior à data de referência (vencida e ativa).'),
    ('Situação', 'Planejar revisão', 'Aprovado com vencimento em até 12 meses da data de referência.'),
    ('Situação', 'Regular', 'Aprovado com vencimento após 12 meses.'),
    ('Prioridade', '1 a 6', '1 vencida ativa · 2 não encontrada · 3 em revisão · 4 em elaboração · 5 vence em até 90 dias · 6 validar cadastro. Dentro da mesma prioridade, a validade mais próxima vem primeiro.'),
    ('KPI', 'Saúde documental', '(Regular + Planejar revisão) ÷ instruções ativas. Leitura: "de cada 100 instruções em uso, X estão aprovadas e dentro da validade". Qualidade cadastral é medida à parte nesta aba.'),
    ('KPI', 'Vencem em 12 meses', 'Aprovadas com validade entre a data de referência e a mesma data +12 meses.'),
    ('KPI', 'Agenda', 'Mesma regra por mês. Mês de concentração = ≥ 2× a média mensal e no mínimo 3 documentos (calculado, sem valor fixo).'),
    ('Área', 'Sem dupla contagem', 'Área consolidada = união das áreas do código nas origens (ex.: Q4 + AOL → Q4/AOL). Gráfico e matriz somam o total de códigos. O filtro "Q4" inclui Q4/AOL.'),
    ('Responsável', 'Normalização', 'PROPER(TRIM()) apenas na camada analítica; o nome original fica na coluna I da base.'),
    ('Data', 'Referência', "Data de referência = HOJE() em '_AUX DASH'!B3 (pode ser fixada para uma reunião)."),
]
for i, (a, b, c3) in enumerate(rules):
    r = RR + 1 + i
    qs[f'B{r}'] = a; qs[f'C{r}'] = b; mput(qs, f'D{r}:H{r}', c3, F(9), align=AL('left', 'center', True, 1))
    qs[f'B{r}'].font = F(8, True, MUTED); qs[f'C{r}'].font = F(9, True, NAVY)
    qs.merge_cells  # no-op
    qs.row_dimensions[r].height = 30
    for c in qs[f'B{r}:H{r}'][0]: c.border = Border(bottom=side())
qs.sheet_properties.tabColor = BLUE

# ---------- OUTROS CONTROLES ----------
oc = wb.create_sheet('OUTROS CONTROLES')
oc.sheet_view.showGridLines = False
oc.column_dimensions['A'].width = 2.5
for c_, w_ in zip('BCDEFG', [28, 12, 52, 26, 16, 2.5]): oc.column_dimensions[c_].width = w_
box(oc, 'B2:F3', fill=FILL(NAVY))
mput(oc, 'B2:D2', 'OUTROS CONTROLES DOCUMENTAIS — LCQ RJ', F(15, True, WHITE), align=AL('left', 'bottom', ind=1))
mput(oc, 'B3:D3', 'Famílias com estrutura própria — mantidas separadas e não somadas ao KPI de Instruções', F(9, color='C9D3E6'), align=AL('left', 'top', ind=1))
c = mput(oc, 'E2:F3', None, F(9, True, WHITE), align=AL('right', 'center', ind=1)); link(c, DB, '← Dashboard'); c.font = F(9, True, WHITE)
oc.row_dimensions[2].height = 28; oc.row_dimensions[3].height = 20
desc = {'NDocs FMG e ANX': 'Formulários, anexos e documentos correlatos NDocs', 'Documentos Transversais': 'Procedimentos, políticas e documentos corporativos aplicáveis ao LCQ',
        'LPP': 'Lições Ponto a Ponto', 'AST AOL': 'Análises de Segurança da Tarefa — analisadores em linha', 'AST Geral': 'Análises de Segurança da Tarefa — gerais',
        "RT's": 'Relatórios técnicos', 'MO': 'Manuais de operação de utilidades'}
for j, h in enumerate(['Família', 'Registros', 'Conteúdo', 'Situação da base', 'Acesso']):
    c = oc.cell(5, 2 + j, h); c.font = F(8, True, WHITE); c.fill = FILL(BLUE); c.alignment = AL('left', 'center', ind=1)
oc.row_dimensions[5].height = 24
for i, (n, c_, r0) in enumerate(fams):
    r = 6 + i; kref, eref = famref(n, c_, r0)
    oc[f'B{r}'] = n; oc[f'C{r}'] = f'=COUNTA({kref})'; oc[f'D{r}'] = desc[n]
    oc[f'E{r}'] = f'=IF(SUMPRODUCT(--ISERROR({eref}))>0,"⚠ "&SUMPRODUCT(--ISERROR({eref}))&" erro(s) legado(s)","✓ Sem erros de fórmula")'
    link(oc[f'F{r}'], n, 'Abrir aba →')
    for c in oc[f'B{r}:F{r}'][0]: c.font = F(10, c.column == 2); c.alignment = AL('left', 'center', ind=1); c.border = Border(bottom=side())
    oc[f'C{r}'].font = F(14, True, NAVY); oc[f'F{r}'].font = F(9, True, BLUE)
    oc.row_dimensions[r].height = 30
oc.conditional_formatting.add('E6:E12', FormulaRule(formula=['LEFT(E6,1)="⚠"'], font=Font(color=AMB, bold=True)))
oc.conditional_formatting.add('E6:E12', FormulaRule(formula=['LEFT(E6,1)="✓"'], font=Font(color=GRN)))
mput(oc, 'B14:F14', 'Registros = linhas preenchidas na coluna-chave de cada aba (a partir da linha 3). Erros legados não foram corrigidos por falta de fonte documental segura — ver QUALIDADE BASE.',
     F(8, color=MUTED, italic=True), align=AL('left', 'center', True))
oc.row_dimensions[14].height = 28
oc.sheet_properties.tabColor = '5B6B85'

# ---------- INÍCIO ----------
ini = wb.create_sheet('INÍCIO')
ini.sheet_view.showGridLines = False
ini.column_dimensions['A'].width = 3
for i in range(2, 14): ini.column_dimensions[L(i)].width = 11
for r in range(1, 40):
    for ci in range(1, 15): ini.cell(r, ci).fill = FILL(BG)
box(ini, 'B2:M5', fill=FILL(NAVY))
mput(ini, 'B2:M3', 'GESTÃO DOCUMENTAL — LCQ RJ', F(22, True, WHITE), align=AL('left', 'bottom', ind=1))
mput(ini, 'B4:M4', 'Cockpit Executivo de Instruções · fonte NDocs Rev.28', F(11, color='C9D3E6'), align=AL('left', 'top', ind=1))
c = mput(ini, 'B5:M5', '="Data de referência: "&TEXT(DataRef,"dd/mm/yyyy")', F(9, color='9FB0CC'), align=AL('left', 'top', ind=1))
for r, h in {2: 26, 3: 22, 4: 20, 5: 18, 6: 18}.items(): ini.row_dimensions[r].height = h
def tile(r, c0, w, title, sub, sheet, accent=BLUE, kpi=None):
    a = L(c0); b = L(c0 + w - 1)
    for rr in range(r, r + 3):
        for ci in range(c0, c0 + w):
            ini.cell(rr, ci).fill = FILL(WHITE)
            ini.cell(rr, ci).border = Border(left=side(accent, 'thick') if ci == c0 else side(BG, 'thick') if False else None,
                                             right=side(BG, 'thick') if ci == c0 + w - 1 else None)
    c = mput(ini, f'{a}{r}:{b}{r}', None, F(11, True, NAVY), align=AL('left', 'bottom', ind=1)); link(c, sheet, title); c.font = F(11, True, NAVY)
    mput(ini, f'{a}{r + 1}:{b}{r + 1}', sub, F(8, color=MUTED), align=AL('left', 'top', True, 1))
    c = mput(ini, f'{a}{r + 2}:{b}{r + 2}', kpi or 'Abrir →', F(9, True, BLUE), align=AL('left', 'center', ind=1))
    if not kpi: link(c, sheet, 'Abrir →'); c.font = F(9, True, BLUE)
def grp(r, t):
    mput(ini, f'B{r}:M{r}', t, F(9, True, MUTED), align=AL('left', 'bottom'))
    ini.row_dimensions[r].height = 22
grp(7, 'COCKPIT')
tile(8, 2, 6, 'Dashboard Executivo', 'Situação geral, riscos, ação necessária, agenda de vencimentos e carteira', DB, NAVY,
     f'=SUM({col("W")})&" instruções · "&COUNTIF({col("R")},"Ação imediata")&" ação imediata · "&TEXT(IFERROR((COUNTIF({col("R")},"Regular")+COUNTIF({col("R")},"Planejar revisão"))/COUNTIF({col("R")},"<>Inativo"),0),"0%")&" saúde"')
tile(8, 8, 6, 'Pesquisa de Documento', 'Localize por código NDocs, número antigo ou parte do título', 'PESQUISA', NAVY)
grp(11, 'BASE E AUDITORIA')
tile(12, 2, 6, 'Base Consolidada', 'Um registro por Código NDocs, com origem rastreável e camada analítica', 'BASE CONSOLIDADA', BLUE)
tile(12, 8, 6, 'Qualidade da Base', 'Duplicidades, divergências, campos vazios, erros legados e regras', 'QUALIDADE BASE', BLUE,
     "='QUALIDADE BASE'!E6&\" crítico(s) · \"&'QUALIDADE BASE'!F6&\" atenção · \"&'QUALIDADE BASE'!G6&\" informativo(s)\"")
grp(15, 'INSTRUÇÕES POR ÁREA (ORIGEM)')
for i, (n, t, s2) in enumerate([('NDocs Q4', 'Q4', 'Instruções da área Q4'), ('NDocs PE_PP', 'PE / PP', 'Polietileno e polipropileno'),
                                ('NDocs AOL', 'AOL', 'Analisadores em linha'), ('NDocs COMUNS', 'Comuns', 'Instruções comuns ao LCQ')]):
    tile(16, 2 + i * 3, 3, t, s2, n, '5B6B85')
grp(19, 'OUTROS CONTROLES DOCUMENTAIS')
tile(20, 2, 6, 'Outros Controles', 'FMG/ANX, Transversais, LPP, AST, RT e MO — separados dos KPIs de instruções', 'OUTROS CONTROLES', '5B6B85')
tile(20, 8, 6, 'Histórico (não usado nos KPIs)', 'Abas NDocs, NDocs (2) e NDocs Salvação — versões anteriores', 'NDocs', GRY)
for r in (8, 12, 16, 20): ini.row_dimensions[r].height = 24; ini.row_dimensions[r + 1].height = 28; ini.row_dimensions[r + 2].height = 20
mput(ini, 'B24:M24', 'Sem macros · filtros por lista suspensa no Dashboard · dados originais preservados nas abas de origem',
     F(8, color=MUTED, italic=True), align=AL('left'))
ini.sheet_properties.tabColor = NAVY

# ---------- abas de origem: padronização visual sem alterar valores ----------
def style_source(ws, kind):
    hdr = 2
    ncol = max((c.column for c in ws[hdr] if c.value not in (None, '')), default=1)
    for r in range(1, ws.max_row + 1):
        for ci in range(1, ncol + 1):
            c = ws.cell(r, ci)
            f0 = c.font
            c.font = Font(name=FONT, size=10 if r > 2 else c.font.sz, bold=f0.b, italic=f0.i, color=f0.color, underline=f0.u)
            if isinstance(c.value, (dt.datetime, dt.date)): c.number_format = 'dd/mm/yyyy'
    ws['A1'].font = F(14, True, NAVY); ws['A1'].alignment = AL('left', 'center')
    ws.row_dimensions[1].height = 30
    for ci in range(1, ncol + 1):
        c = ws.cell(hdr, ci); c.font = F(9, True, WHITE); c.fill = FILL(NAVY if kind != 'hist' else GRY)
        c.alignment = AL('left', 'center', True, 1)
    ws.row_dimensions[hdr].height = 30
    for ci in range(1, ncol + 1):
        mx = max((len(str(ws.cell(r, ci).value)) for r in range(2, min(ws.max_row, 300) + 1) if ws.cell(r, ci).value is not None), default=8)
        if any(isinstance(ws.cell(r, ci).value, dt.datetime) for r in range(3, min(ws.max_row, 50))): mx = 11
        ws.column_dimensions[L(ci)].width = max(9, min(mx + 2, 55))
    ws.freeze_panes = 'A3'
    if not ws.tables and not ws.auto_filter.ref:
        ws.auto_filter.ref = f'A{hdr}:{L(ncol)}{ws.max_row}'
    ws.sheet_view.showGridLines = False
    for r in range(3, ws.max_row + 1):
        for ci in range(1, ncol + 1):
            c = ws.cell(r, ci)
            if not (c.border.bottom and c.border.bottom.style): c.border = Border(bottom=side())
            if c.alignment.wrap_text is None: c.alignment = AL('left', 'center', ci in (4, 5) and kind != 'hist')
    # navegação (substitui o antigo texto "DASHBOARD PREMIUM" em M1)
    for c in ws[1]:
        if c.value == 'DASHBOARD PREMIUM': c.value = None
    a = max([ncol] + [m.max_col for m in ws.merged_cells.ranges if m.min_row <= 1]) + 2
    link(ws.cell(1, a), 'INÍCIO', '← Início'); ws.cell(1, a).font = F(9, True, BLUE)
    link(ws.cell(1, a + 1), DB, 'Dashboard'); ws.cell(1, a + 1).font = F(9, True, BLUE)
    ws.column_dimensions[L(a)].width = max(ws.column_dimensions[L(a)].width or 0, 10)
    ws.column_dimensions[L(a + 1)].width = max(ws.column_dimensions[L(a + 1)].width or 0, 11)
    ws.sheet_properties.tabColor = {'instr': '5B6B85', 'outro': '8A94A3', 'hist': 'C5CBD3'}[kind]
    if kind == 'hist':
        ws.cell(1, a + 3, 'HISTÓRICO — versão anterior; não usada nos KPIs').font = F(9, True, GRY, italic=True)

for n in ['NDocs Q4', 'NDocs PE_PP', 'NDocs AOL', 'NDocs COMUNS']: style_source(wb[n], 'instr')
for n in ['NDocs FMG e ANX', 'Documentos Transversais', 'LPP', 'AST AOL', 'AST Geral', "RT's", 'MO']: style_source(wb[n], 'outro')
for n in ['NDocs', 'NDocs (2)', 'NDocs Salvação']: style_source(wb[n], 'hist')
# destaque profissional dos erros legados
for rng in ['LPP!F100:F124', "'AST Geral'!E42"]:
    sh, rr = rng.rsplit('!', 1); sh = sh.strip("'")
    from openpyxl.utils.cell import range_boundaries
    c1, r1_, c2, r2_ = range_boundaries(rr)
    for row in wb[sh].iter_rows(min_row=r1_, max_row=r2_, min_col=c1, max_col=c2):
        for c in row:
            c.fill = FILL(AMBL); c.font = F(10, False, AMB)
wb['LPP'].cell(1, 18, 'F100:F124 = erro legado (REF) preservado — ver QUALIDADE BASE').font = F(8, color=AMB, italic=True)

bs.sheet_properties.tabColor = BLUE
bs.row_dimensions[3].height = 16

# ---------- ordem das abas ----------
order_ = ['INÍCIO', DB, 'PESQUISA', 'BASE CONSOLIDADA', 'QUALIDADE BASE', 'NDocs Q4', 'NDocs PE_PP', 'NDocs AOL', 'NDocs COMUNS',
          'OUTROS CONTROLES', 'NDocs FMG e ANX', 'Documentos Transversais', 'LPP', 'AST AOL', 'AST Geral', "RT's", 'MO',
          'NDocs', 'NDocs (2)', 'NDocs Salvação', '_AUX DASH']
wb._sheets = [wb[n] for n in order_]
wb.active = 0
for ws in wb: ws.sheet_view.tabSelected = ws.title == 'INÍCIO'
for ws_ in [ini, ps, qs, oc, bs, ax]:
    ws_.page_setup.orientation = 'landscape'; ws_.page_setup.fitToWidth = 1; ws_.page_setup.fitToHeight = 0
    ws_.sheet_properties.pageSetUpPr.fitToPage = True
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('ok', OUT, 'achados', len(fd), collections.Counter(x[0] for x in fd))
