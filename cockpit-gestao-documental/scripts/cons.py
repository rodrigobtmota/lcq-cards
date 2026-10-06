import openpyxl, collections, datetime as dt, re
REF=dt.datetime(2026,10,6)
SRC=['NDocs Q4','NDocs PE_PP','NDocs AOL','NDocs COMUNS']
def pdate(v):
    if isinstance(v,dt.datetime): return v
    if isinstance(v,str):
        m=re.search(r'(\d{2})/(\d{2})/(\d{4})',v)
        if m: return dt.datetime(int(m[3]),int(m[2]),int(m[1]))
    return None
def nstatus(s):
    s=(s or '').replace('\xa0',' ').strip()
    return {'Avaiação de uso':'Avaliação de uso','Cancelada':'Cancelado'}.get(s,s)
def load(path='sp.xlsx'):
    wb=openpyxl.load_workbook(path,data_only=True)
    recs=[]; lvs=[]
    for n in SRC:
        ws=wb[n]
        for r in range(3,ws.max_row+1):
            v=[ws.cell(r,c).value for c in range(1,13)]
            if not v[1]: continue
            d=dict(sheet=n,row=r,tipo=str(v[0]).strip(),cod=str(v[1]).strip(),ver=v[2],ant=v[3],tit=v[4],st=v[5],val=v[6],area=v[7],abr=v[8],resp=v[9],norma=v[10],com=v[11])
            (recs if d['tipo']=='Instrução' else lvs).append(d)
    g=collections.OrderedDict()
    for d in recs: g.setdefault(d['cod'],[]).append(d)
    return wb,recs,lvs,g
def area_union(l):
    toks=[]
    for d in l:
        for t in str(d['area']).strip().split('/'):
            if t not in toks: toks.append(t)
    order=['Q4','PE','PP','AOL','COMUM']
    toks=sorted(toks,key=lambda t: order.index(t) if t in order else 9)
    return '/'.join(toks)
def sit(d):
    s=nstatus(d['st']); v=pdate(d['val'])
    if s in ('Cancelado','Fora de Uso'): return 'Inativo'
    if s in ('Revisão','Em Elaboração'): return 'Em tratamento'
    if s!='Aprovado' or v is None: return 'Validar cadastro'
    days=(v-REF).days
    if days<0: return 'Ação imediata'
    if days<=365: return 'Planejar revisão'
    return 'Regular'
if __name__=='__main__':
    wb,recs,lvs,g=load()
    P=[ (k,l[0]) for k,l in g.items()]
    print('linhas instr',len(recs),'lvs',len(lvs),'unicos',len(g))
    print(collections.Counter(nstatus(d['st']) for k,d in P))
    print(collections.Counter(sit(d) for k,d in P))
    print(collections.Counter(area_union(l) for l in g.values()))
    for k,d in P:
        s=sit(d)
        if s in('Ação imediata','Em tratamento','Validar cadastro','Inativo'): print(s,k,nstatus(d['st']),d['val'],area_union(g[k]),d['resp'],str(d['tit'])[:40])
    m=collections.Counter()
    for k,d in P:
        v=pdate(d['val'])
        if v and sit(d)!='Inativo' and 0<= (v-REF).days and v< dt.datetime(2027,10,1): m[(v.year,v.month)]+=1
    print(sorted(m.items()))
    print('<=90',sum(1 for k,d in P if sit(d)=='Planejar revisão' and (pdate(d['val'])-REF).days<=90))
